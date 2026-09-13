"""Detección de tipo por contenido, límites y extracción por formato (§5.8, §8.7, §8.1).

La prueba central del archivo es la última: **un PDF escaneado sin capa de texto se
rechaza con el mensaje del contrato** (`DOCUMENT_UNREADABLE`, §8.1), no con una traza ni
con un `500`.
"""

from __future__ import annotations

import pytest
from tests.ingestion.conftest import docx_minimo, pdf_con_texto, pdf_en_blanco, texto_largo

from app.core.errors import DEFAULT_MESSAGES
from app.models.enums import DocumentType
from app.modules.ingestion import extraccion

# ---------------------------------------------------------------------------
# Tipo real por contenido, nunca por extensión (§8.7)
# ---------------------------------------------------------------------------


def test_detecta_pdf_aunque_la_extension_mienta() -> None:
    """Un `.txt` cuyo contenido empieza por `%PDF-` es un PDF, y se anota la incoherencia."""
    datos = pdf_con_texto(["Hola mundo"])
    detectado = extraccion.detectar_tipo("apuntes.txt", datos)
    assert detectado.document_type == DocumentType.PDF
    assert detectado.mime == "application/pdf"
    assert detectado.extension_incoherente is True


def test_detecta_docx_por_el_contenido_del_zip() -> None:
    """Un DOCX se reconoce porque el ZIP contiene `word/document.xml`."""
    detectado = extraccion.detectar_tipo("material.bin", docx_minimo(["Un párrafo."]))
    assert detectado.document_type == DocumentType.DOCX


def test_zip_que_no_es_docx_se_rechaza_con_415() -> None:
    """Un ZIP cualquiera no es material admitido."""
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as paquete:
        paquete.writestr("hola.txt", "contenido")
    with pytest.raises(extraccion.TipoNoAdmitido) as excepcion:
        extraccion.detectar_tipo("material.docx", buffer.getvalue())
    assert excepcion.value.status_code == 415
    assert excepcion.value.code == "UNSUPPORTED_FILE_TYPE"


def test_ejecutable_disfrazado_de_pdf_se_rechaza() -> None:
    """El nombre dice `.pdf`, los magic bytes dicen ejecutable: manda el contenido."""
    with pytest.raises(extraccion.TipoNoAdmitido) as excepcion:
        extraccion.detectar_tipo("manual.pdf", b"MZ\x90\x00" + b"\x00" * 200)
    assert excepcion.value.status_code == 415
    assert "Ejecutable" in str(excepcion.value.details.get("detected", ""))


def test_markdown_se_distingue_del_texto_plano() -> None:
    """Encabezados y vallas de código marcan Markdown; el texto liso es TXT."""
    markdown = b"# Titulo\n\nUn parrafo.\n\n```sql\nSELECT 1;\n```\n\n- punto\n"
    assert extraccion.detectar_tipo("a.dat", markdown).document_type == DocumentType.MARKDOWN
    assert extraccion.detectar_tipo("a.dat", b"Solo texto corrido sin marcas.").document_type == (
        DocumentType.TXT
    )


def test_archivo_vacio_se_rechaza() -> None:
    """Un archivo de cero bytes no es ningún formato."""
    with pytest.raises(extraccion.TipoNoAdmitido):
        extraccion.detectar_tipo("vacio.txt", b"")


# ---------------------------------------------------------------------------
# Límites de §5.8
# ---------------------------------------------------------------------------


def test_tamano_maximo_lanza_413() -> None:
    """`ingestion.max_file_mb` se comprueba antes de tocar el disco."""
    with pytest.raises(extraccion.ArchivoDemasiadoGrande) as excepcion:
        extraccion.validar_tamano(b"x" * 2048, max_bytes=1024)
    assert excepcion.value.status_code == 413
    assert excepcion.value.code == "FILE_TOO_LARGE"


def test_demasiadas_paginas_lanza_413() -> None:
    """`ingestion.max_pdf_pages` acota el coste de la ingesta."""
    with pytest.raises(extraccion.ArchivoDemasiadoGrande) as excepcion:
        extraccion.extraer(
            pdf_en_blanco(paginas=4),
            DocumentType.PDF,
            max_paginas=2,
            min_chars_por_pagina=50,
        )
    assert excepcion.value.details["page_count"] == 4


def test_material_insuficiente_lleva_min_words_en_details() -> None:
    """`DOCUMENT_UNREADABLE` con `details.min_words`, tal como pide §8.1."""
    extraido = extraccion.extraer(
        b"Cuatro palabras nada mas", DocumentType.TXT, max_paginas=300, min_chars_por_pagina=50
    )
    with pytest.raises(extraccion.DocumentoIlegible) as excepcion:
        extraccion.validar_material_suficiente(extraido, min_palabras=300)
    assert excepcion.value.details["min_words"] == 300
    assert excepcion.value.details["reason"] == "insufficient_content"


# ---------------------------------------------------------------------------
# Limpieza y métricas
# ---------------------------------------------------------------------------


def test_limpiar_texto_conserva_la_sangria_del_codigo() -> None:
    """La sangría es significativa: se borran los espacios finales, nunca los iniciales."""
    crudo = "def f():\r\n    return 1   \r\n\r\n\r\n\r\nfin"
    limpio = extraccion.limpiar_texto(crudo)
    assert "    return 1" in limpio
    assert "   \n" not in limpio
    assert "\n\n\n" not in limpio


def test_detecta_idioma_espanol_e_ingles() -> None:
    """Heurística de palabras vacías; ante la duda, español (el producto es en español)."""
    assert extraccion.detectar_idioma("el gato de la casa y el perro en el patio") == "es"
    assert extraccion.detectar_idioma("the cat of the house and the dog in the yard") == "en"
    assert extraccion.detectar_idioma("") == "es"


def test_extrae_pdf_con_texto_y_cuenta_paginas() -> None:
    """Un PDF con capa de texto devuelve su contenido, su página y sus métricas."""
    extraido = extraccion.extraer(
        pdf_con_texto(["Capitulo 1. JOINs", texto_largo(400)]),
        DocumentType.PDF,
        max_paginas=300,
        min_chars_por_pagina=50,
    )
    assert "JOINs" in extraido.texto
    assert extraido.page_count == 1
    assert extraido.word_count > 300
    assert extraido.pagina_de(5) == 1


def test_extrae_docx_con_tabla_como_markdown() -> None:
    """Las tablas de Word salen como tabla de tuberías: bloque atómico para el troceador."""
    import io

    import docx

    documento = docx.Document()
    documento.add_heading("Comparativa", level=1)
    tabla = documento.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Tipo"
    tabla.cell(0, 1).text = "Uso"
    tabla.cell(1, 0).text = "LEFT JOIN"
    tabla.cell(1, 1).text = "Conserva la izquierda"
    buffer = io.BytesIO()
    documento.save(buffer)

    extraido = extraccion.extraer(
        buffer.getvalue(), DocumentType.DOCX, max_paginas=300, min_chars_por_pagina=50
    )
    assert "| Tipo | Uso |" in extraido.texto
    assert "| --- | --- |" in extraido.texto
    assert "| LEFT JOIN | Conserva la izquierda |" in extraido.texto


# ---------------------------------------------------------------------------
# LA prueba del contrato: PDF escaneado sin texto
# ---------------------------------------------------------------------------


def test_pdf_escaneado_sin_texto_se_rechaza_con_el_mensaje_del_contrato() -> None:
    """§8.1: `422 DOCUMENT_UNREADABLE` con el mensaje del catálogo y `reason` en `details`.

    El mensaje es el que la app muestra **tal cual**; el motivo técnico
    (`scanned_pdf_no_text`) viaja en `details` y es el mismo que lleva el evento
    `DOCUMENT_FAILED` de §4.2.
    """
    with pytest.raises(extraccion.DocumentoIlegible) as excepcion:
        extraccion.extraer(
            pdf_en_blanco(paginas=3),
            DocumentType.PDF,
            max_paginas=300,
            min_chars_por_pagina=50,
        )
    error = excepcion.value
    assert error.status_code == 422
    assert error.code == "DOCUMENT_UNREADABLE"
    assert error.message == DEFAULT_MESSAGES["DOCUMENT_UNREADABLE"]
    assert error.message == "No pudimos leer el documento. Prueba con un archivo de texto real."
    assert error.details["reason"] == "scanned_pdf_no_text"
    assert error.details["page_count"] == 3
    assert error.details["min_chars_per_page"] == 50
